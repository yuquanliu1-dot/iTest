#!/usr/bin/env python3
"""把 dump_cases_payload 落盘的用例 payload 渲染为 XMind/Word 文件，按平台产物协议输出下载清单。

用法（由 AI 经 run_skill_script 调用，cwd = 技能根）：
  export_cases.py '{"payload":"<绝对路径>","format":"xmind|docx","title":"可选根标题"}'

payload 文件由 dump_cases_payload 工具生成：{baseName, count, tree}，tree 节点 {title, children?}，
层级 root → [版本(跨版本时)] → 分组 → 用例 → 预置条件/测试步骤/预期结果（叶子是已渲染好的文本行）。
层级与排序在 TS 侧定死，本脚本只按 payload 顺序渲染。

- xmind：content.json(XMind 2020+/ZEN) + content.xml(XMind 8) 双格式 + metadata + manifest，zipfile 打包。
- docx：表格式——分组做章节标题，每用例一张 2 列表（用例名称/所属分组/预置条件/测试步骤/预期结果，
  步骤多行=单元格内多段落）；优先 python-docx，ImportError 降级标准库 zipfile 手写 OOXML。

产物协议：文件写到 .out/ 下，stdout 最后一行打印
  __ITEST_ARTIFACTS__ [{"name":"…","path":".out/…","count":N}]
平台据此落库并推送下载卡，用户点击即下载。
"""
import json
import os
import re
import sys
import uuid
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

OUT_DIR = ".out"
SECTION_TITLES = ("预置条件", "测试步骤", "预期结果")
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def fail(msg, code=2):
    print(f"错误：{msg}", file=sys.stderr)
    sys.exit(code)


def load_args():
    if len(sys.argv) < 2:
        fail("缺少参数 JSON（usage: export_cases.py '{\"payload\":…,\"format\":…}'）")
    try:
        data = json.loads(sys.argv[1])
    except json.JSONDecodeError as e:
        fail(f"参数 JSON 解析失败：{e}")
    if not isinstance(data, dict):
        fail("参数顶层必须是对象")
    payload_path = str(data.get("payload") or "").strip()
    if not payload_path:
        fail("缺少 payload（先调 dump_cases_payload 拿 payloadPath）")
    fmt = str(data.get("format") or "").strip().lower()
    if fmt not in ("xmind", "docx"):
        fail(f"format 必须是 xmind 或 docx，收到：{fmt or '(空)'}")
    title = str(data.get("title") or "").strip()
    return payload_path, fmt, title


def read_payload(path, title_override):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except OSError as e:
        fail(f"payload 文件不可读：{path}（{e}；先调 dump_cases_payload 重新生成）")
    except json.JSONDecodeError as e:
        fail(f"payload JSON 解析失败：{e}")
    tree = data.get("tree")
    if not isinstance(tree, dict) or not str(tree.get("title") or "").strip():
        fail("payload 缺少有效 tree（重新 dump）")
    base = title_override or str(data.get("baseName") or "").strip() or str(tree["title"])
    if title_override:
        tree = {**tree, "title": title_override}  # 根标题同步覆盖（xmind 根节点/docx 文档标题）
    count = data.get("count")
    count = count if isinstance(count, int) and count >= 0 else 0
    return base, count, tree


def out_filename(base, ext):
    stem = re.sub(r'[\\/:*?"<>|\s]+', "-", base).strip("-") or "cases"
    return f"{stem}-用例-{datetime.now():%Y%m%d-%H%M%S}.{ext}"


# ---------- 结构解释（docx 用）：root → [版本] → 分组 → 用例 → 三要素 ----------

def is_case_node(node):
    children = node.get("children") or []
    if not children:
        return False
    return any(str(c.get("title")) in SECTION_TITLES for c in children)


def walk_sections(tree):
    """→ [(版本标题|None, [分组节点…])]，保持 payload 顺序。"""
    out = []
    for lvl1 in tree.get("children") or []:
        children = lvl1.get("children") or []
        if children and all(is_case_node(c) for c in children):
            out.append((None, [lvl1]))
        elif children:
            out.append((str(lvl1.get("title")), children))
    return out


def case_rows(case, group_title):
    """用例节点 → [(标签, [文本行…])]，空字段行跳过。"""
    rows = [("用例名称", [str(case.get("title"))]), ("所属分组", [group_title])]
    by_title = {}
    for c in case.get("children") or []:
        by_title[str(c.get("title"))] = [str(leaf.get("title")) for leaf in (c.get("children") or [])]
    for label in SECTION_TITLES:
        lines = [l for l in by_title.get(label, []) if l.strip()]
        if lines:
            rows.append((label, lines))
    return rows


# ---------- xmind 引擎（移植自 services/xmindExport.service.ts） ----------

def topic_to_json(n, is_root):
    t = {"id": str(uuid.uuid4()), "class": "topic", "title": n["title"]}
    if is_root:
        t["structureClass"] = "org.xmind.ui.logic.right"
    children = n.get("children")
    if children:
        t["children"] = {"attached": [topic_to_json(c, False) for c in children]}
    return t


def topic_to_xml(n, is_root):
    attrs = ' structure-class="org.xmind.ui.logic.right"' if is_root else ""
    inner = f"<title>{escape(str(n['title']))}</title>"
    children = n.get("children")
    kids = ""
    if children:
        inner_topics = "".join(topic_to_xml(c, False) for c in children)
        kids = f'<children><topics type="attached">{inner_topics}</topics></children>'
    return f'<topic id="{uuid.uuid4()}"{attrs}>{inner}{kids}</topic>'


def build_xmind(tree):
    content_json = json.dumps(
        [{"id": str(uuid.uuid4()), "class": "sheet", "title": "Sheet 1", "rootTopic": topic_to_json(tree, True)}],
        ensure_ascii=False, separators=(",", ":"),
    )
    content_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
        '<xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0" '
        'xmlns:xlink="http://www.w3.org/1999/xlink" version="2.0">\n'
        f'  <sheet id="{uuid.uuid4()}">\n'
        f"    {topic_to_xml(tree, True)}\n"
        "  </sheet>\n"
        "</xmap-content>\n"
    )
    metadata_json = json.dumps(
        {"dataStructureVersion": "2", "layoutEngineVersion": "3", "creator": {"name": "iTest", "version": "1.0"}},
        ensure_ascii=False, separators=(",", ":"),
    )
    manifest_json = json.dumps(
        {"file-entries": {"content.json": {}, "content.xml": {}, "metadata.json": {}}},
        ensure_ascii=False, separators=(",", ":"),
    )
    return {
        "content.json": content_json,
        "content.xml": content_xml,
        "metadata.json": metadata_json,
        "manifest.json": manifest_json,
    }


# ---------- docx 引擎：python-docx 优先 ----------

def write_docx_with_python(path, base, sections):
    from docx import Document

    doc = Document()
    doc.add_heading(base, 0)
    for version_title, groups in sections:
        if version_title:
            doc.add_heading(version_title, 1)
        for group in groups:
            doc.add_heading(str(group.get("title")), 2 if version_title else 1)
            for case in group.get("children") or []:
                table = doc.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                for label, lines in case_rows(case, str(group.get("title"))):
                    cells = table.add_row().cells
                    cells[0].text = label
                    cells[1].text = lines[0]
                    for extra in lines[1:]:
                        cells[1].add_paragraph(extra)
                doc.add_paragraph("")
    doc.save(path)


# ---------- docx 引擎：标准库 OOXML 降级（骨架同 word-test-report） ----------

def w_paragraph(text, style=None, bold=False):
    pr = ""
    if style:
        pr += f'<w:pStyle w:val="{style}"/>'
    run_pr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return (
        f'<w:p>{pr}<w:r>{run_pr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
    )


def w_kv_table(rows):
    """2 列 label/value 表；value 单元格多行=多段落，label 加粗。"""
    borders = (
        "<w:tblBorders>"
        + "".join(
            f'<w:{side} w:val="single" w:sz="4" w:color="999999"/>'
            for side in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        + "</w:tblBorders>"
    )
    rows_xml = []
    for label, lines in rows:
        label_cell = f'<w:tc><w:tcPr/><w:p><w:r><w:rPr><w:b/></w:rPr>' \
                     f'<w:t xml:space="preserve">{escape(label)}</w:t></w:r></w:p></w:tc>'
        value_paras = "".join(w_paragraph(l) for l in lines)
        value_cell = f"<w:tc><w:tcPr/>{value_paras}</w:tc>"
        rows_xml.append(f"<w:tr>{label_cell}{value_cell}</w:tr>")
    return (
        f'<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/>{borders}</w:tblPr>'
        + "".join(rows_xml)
        + "</w:tbl>"
    )


def write_docx_with_zipfile(path, base, sections):
    body = [w_paragraph(base, style="Title")]
    for version_title, groups in sections:
        if version_title:
            body.append(w_paragraph(version_title, style="Heading1"))
        for group in groups:
            body.append(w_paragraph(str(group.get("title")), style="Heading2" if version_title else "Heading1"))
            for case in group.get("children") or []:
                body.append(w_kv_table(case_rows(case, str(group.get("title")))))
                body.append(w_paragraph(""))
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W_NS}"><w:body>'
        + "".join(body)
        + '<w:sectPr/></w:body></w:document>'
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        "</Relationships>"
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:styles xmlns:w="{W_NS}">'
        '<w:style w:type="paragraph" w:styleId="Normal" w:default="1"><w:name w:val="Normal"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/>'
        '<w:pPr><w:jc w:val="center"/></w:pPr><w:rPr><w:b/><w:sz w:val="56"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/>'
        '<w:rPr><w:b/><w:sz w:val="32"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/>'
        '<w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style>'
        "</w:styles>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/_rels/document.xml.rels", doc_rels)
        z.writestr("word/styles.xml", styles)


def main():
    payload_path, fmt, title = load_args()
    base, count, tree = read_payload(payload_path, title)

    os.makedirs(OUT_DIR, exist_ok=True)
    fname = out_filename(base, "xmind" if fmt == "xmind" else "docx")
    path = os.path.join(OUT_DIR, fname)

    if fmt == "xmind":
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, content in build_xmind(tree).items():
                z.writestr(name, content)
        engine = "xmind-zip"
    else:
        sections = walk_sections(tree)
        if not sections:
            fail("payload 里没有可导出的分组/用例")
        try:
            write_docx_with_python(path, base, sections)
            engine = "python-docx"
        except ImportError:
            write_docx_with_zipfile(path, base, sections)
            engine = "stdlib-ooxml"

    size = os.path.getsize(path)
    print(f"已生成 {path}（{size} 字节，{count} 条用例，引擎 {engine}）")
    print(f'__ITEST_ARTIFACTS__ [{{"name":"{fname}","path":"{path}","count":{count}}}]')


if __name__ == "__main__":
    main()
